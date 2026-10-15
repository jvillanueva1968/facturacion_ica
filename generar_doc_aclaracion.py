from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import datetime

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

ICA_VERDE = RGBColor(0x06, 0x84, 0x60)


def add_table(headers, rows, widths_size=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(widths_size)
    for row_idx, row in enumerate(rows, 1):
        for col_idx, value in enumerate(row):
            table.cell(row_idx, col_idx).text = value
            for p in table.cell(row_idx, col_idx).paragraphs:
                for r in p.runs:
                    r.font.size = Pt(widths_size)
    return table


# Titulo
title = doc.add_heading('SOLICITUD DE ACLARACION PREVIA - CUATRO (4) PUNTOS CLAVE', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run(
    'Para estructurar y tramitar formatos, Acuerdos de Confidencialidad (NDA) o cartas de '
    'autorizacion ante la Direccion de Investigacion y Transferencia y Secretaria General'
)
run.font.size = Pt(12)
run.font.color.rgb = ICA_VERDE

naturaleza = doc.add_paragraph()
naturaleza.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = naturaleza.add_run('Documento elaborado en el marco de un PROYECTO DE GRADO (Opcion de Grado)')
run.font.size = Pt(11)
run.bold = True

doc.add_paragraph()

# Informacion del documento
info_table = doc.add_table(rows=7, cols=2)
info_table.style = 'Light Grid Accent 1'
info_data = [
    ('Proyecto:', 'Proyecto de grado - Validacion OCR de vouchers/comprobantes de pago ICA'),
    ('Naturaleza:', 'Proyecto de grado (Opcion de Grado) - trabajo academico con entidad receptora'),
    ('Tramites que se preparan:', 'Formatos institucionales, Acuerdo de Confidencialidad (NDA) y/o carta de autorizacion'),
    ('Version:', '1.0'),
    ('Fecha:', datetime.date.today().strftime('%d/%m/%Y')),
    ('Elaborado por:', 'Equipo del proyecto de grado - [Nombre del(as) estudiante(s)] - [Programa / Universidad]'),
    ('Destinatarios:', 'Direccion de Investigacion y Transferencia / Secretaria General - ICA'),
]
for i, (label, value) in enumerate(info_data):
    info_table.cell(i, 0).text = label
    info_table.cell(i, 1).text = value
    for cell in info_table.rows[i].cells:
        for p in cell.paragraphs:
            p.style = doc.styles['Normal']
            for r in p.runs:
                r.font.size = Pt(10)

doc.add_paragraph()

# 1. CONTEXTO
doc.add_heading('1. CONTEXTO Y OBJETIVO', level=2)
doc.add_paragraph(
    'El presente documento corresponde a un proyecto de grado desarrollado por estudiantes del '
    'programa de [Programa Academico] de [Universidad], el cual tiene por objeto tecnico la '
    'extraccion y validacion automatica de informacion contenida en vouchers/comprobantes de pago '
    'mediante tecnicas de Reconocimiento Optico de Caracteres (OCR), en el marco del proceso de '
    'facturacion automatizada del Instituto Colombiano Agropecuario (ICA) - SNRI + SIGMA.'
)
doc.add_paragraph(
    'Con el fin de cumplir con los requisitos formales de la Direccion de Investigacion y '
    'Transferencia y de la Secretaria General, y para poder estructurar y tramitar los formatos, '
    'los Acuerdos de Confidencialidad (NDA) o las cartas de autorizacion que resulten requeridos, '
    'le solicitamos cordialmente su aclaracion sobre los siguientes cuatro (4) puntos clave:'
)

# 2. LOS CUATRO PUNTOS CLAVE
doc.add_heading('2. PUNTOS CLAVE QUE REQUIEREN ACLARACION', level=2)

puntos = [
    (
        '2.1 Origen y Vinculacion Institucional de los Datos',
        'La informacion de los vouchers/comprobantes es suministrada directamente por el ICA '
        'mediante algun enlace o contacto oficial, o proviene de un tercero/proveedor vinculado? '
        'Indiquenos si existe un interlocutor o representante legal en la entidad para la firma '
        'de los acuerdos.',
        [
            ('Fuente de los datos', 'ICA (canal oficial) / Tercero-proveedor vinculado / Otro'),
            ('Enlace o contacto oficial', 'Correo, area o funcional que remite la informacion'),
            ('Representante legal para firmas', 'Nombre, cargo y area facultada para suscribir acuerdos'),
        ],
    ),
    (
        '2.2 Estado y Grado de Anonimizacion de la Informacion',
        'Los vouchers contienen datos personales identificables de personas naturales (nombres, '
        'cedulas, numeros de cuenta o tarjeta) o informacion estrictamente comercial/institucional? '
        'La base de datos sera entregada previamente anonimizada por el ICA o se requiere que el '
        'equipo ejecute el protocolo de anonimizacion como primer paso?',
        [
            ('Contenido de los datos', 'Datos personales identificables / Comercial-institucional / Mixto'),
            ('Estado al momento de la entrega', 'Entrega anonimizada por el ICA / Anonimizacion a cargo del equipo'),
            ('Protocolo de anonimizacion', 'Guia o formato institucional aplicable (si existe)'),
        ],
    ),
    (
        '2.3 Plan de Custodia y Almacenamiento',
        'En que infraestructura se tiene previsto realizar el almacenamiento y procesamiento de los '
        'datos (servidores de U_Compensar, repositorios locales o servicios en la nube)?',
        [
            ('Infraestructura prevista', 'Servidores de U_Compensar / Repositorios locales / Nube'),
            ('Requisitos de custodia', 'Restricciones, accesos y tiempos de retencion exigidos'),
            ('Responsable de custodia', 'Area o persona que validara el cumplimiento'),
        ],
    ),
    (
        '2.4 Tratamiento de los Entregables y Publicacion',
        'Al finalizar el trabajo de grado, se tiene previsto publicar la monografia/codigo en el '
        'Repositorio CRAI institucional bajo acceso abierto, o la entidad exigira reserva/'
        'confidencialidad temporal sobre los datos y los resultados analiticos?',
        [
            ('Publicacion de entregables', 'Repositorio CRAI (acceso abierto) / Reserva solicitada'),
            ('Alcance de la reserva', 'Datos, resultados analiticos, codigo fuente o todos'),
            ('Duracion de la confidencialidad', 'Periodo exigido de reserva temporal'),
        ],
    ),
]

for titulo, pregunta, campos in puntos:
    doc.add_heading(titulo, level=3)
    p = doc.add_paragraph()
    p.add_run('Pregunta: ').bold = True
    p.add_run(pregunta)
    add_table(['Aspecto por aclarar', 'Respuesta ICA (por diligenciar)'],
              [(aspecto, '') for aspecto, _ in campos])
    doc.add_paragraph()

# 3. TRAMITES QUE DEPENDEN DE LAS RESPUESTAS
doc.add_heading('3. TRAMITES QUE DEPENDEN DE ESTAS RESPUESTAS', level=2)
doc.add_paragraph(
    'Con base en las respuestas a los cuatro puntos anteriores se estructuraran y radicaran los '
    'siguientes documentos ante las instancias correspondientes:'
)
add_table(
    ['Punto', 'Tramite a preparar / radicar', 'Instancia'],
    [
        ('2.1', 'Formatos de solicitud de informacion y carta de autorizacion de uso de datos', 'Direccion de Investigacion y Transferencia'),
        ('2.2', 'Acuerdo de Confidencialidad (NDA) y, de ser necesario, protocolo de anonimizacion', 'Secretaria General / Direccion de Investigacion y Transferencia'),
        ('2.3', 'Anexo de plan de custodia, almacenamiento y seguridad de la informacion', 'Direccion de Investigacion y Transferencia'),
        ('2.4', 'Clausula de publicacion, deposito en Repositorio CRAI y/o reserva temporal de resultados', 'Secretaria General'),
    ],
)

doc.add_paragraph()

# 4. COMPROMISOS DEL EQUIPO
doc.add_heading('4. COMPROMISOS PREVIOS DEL EQUIPO DEL PROYECTO DE GRADO', level=2)
compromisos = [
    'Uso unico y exclusivo de los datos para los fines academicos del proyecto de grado',
    'Aplicacion del protocolo de anonimizacion suministrado por la entidad o aprobado por esta',
    'Custodia de la informacion en la infraestructura que la entidad autorice',
    'No divulgacion de datos personales, resultados ni artefactos sin autorizacion expresa',
    'Devolucion o eliminacion segura de los datos al finalizar el proyecto de grado',
    'Cite obligatorio de la entidad y de los insumos institucionales en los entregables academicos',
]
for c in compromisos:
    doc.add_paragraph(c, style='List Bullet')

doc.add_paragraph()

# 5. ACLARACION ADICIONAL
doc.add_heading('5. ACLARACION ADICIONAL', level=2)
doc.add_paragraph(
    'Se aclara que el presente trabajo tiene naturaleza de PROYECTO DE GRADO (opcion de grado), '
    'no constituye un contrato comercial ni un servicio prestado a la entidad, y que su unico '
    'proposito es academico. Por lo anterior, se solicita que los formatos, acuerdos o '
    'cartas que deban suscribirse contemplen esta condicion y, de ser posible, apliquen las '
    'condiciones institucionales previstas para convenios de practica, pasantia o apoyo a '
    'proyectos academicos.'
)

doc.add_paragraph()
p = doc.add_paragraph('Quedamos atentos a sus comentarios para avanzar con los tramites requeridos.')
p.runs[0].bold = True

doc.add_paragraph()

# Firma
doc.add_paragraph('_' * 60)
doc.add_paragraph('Atentamente,').runs[0].bold = True
doc.add_paragraph('_________________________________')
doc.add_paragraph('[Nombre del(as) estudiante(s)]')
doc.add_paragraph('[Programa Academico] - [Universidad]')
doc.add_paragraph('Proyecto de grado - Validacion OCR de vouchers ICA')
doc.add_paragraph()
doc.add_paragraph('_________________________________')
doc.add_paragraph('Asesor(a) de proyecto de grado')
doc.add_paragraph('[Nombre] - [Programa / Universidad]')

doc.add_paragraph()

# Respuesta institucional
doc.add_heading('DILIGENCIAMIENTO POR LA ENTIDAD', level=2)
add_table(
    ['Item', 'Respuesta de la entidad'],
    [(f'Punto {i}', '') for i in range(1, 5)],
)

doc.add_paragraph()
doc.add_paragraph('_________________________________')
doc.add_paragraph('Direccion de Investigacion y Transferencia - ICA')
doc.add_paragraph()
doc.add_paragraph('_________________________________')
doc.add_paragraph('Secretaria General - ICA')

output_path = 'Solicitud_Aclaracion_4_Puntos_Proyecto_Grado_ICA.docx'
doc.save(output_path)
print(f'Documento guardado: {output_path}')
