from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import datetime

doc = Document()

# Estilos
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Título
title = doc.add_heading('SOLICITUD DE MUESTRAS DE VOUCHERS/COMPROBANTES DE PAGO', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Subtítulo
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Para Entrenamiento y Validacion de OCR - Facturacion Automatizada ICA - SNRI + SIGMA')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x06, 0x84, 0x60)  # ICA Verde

doc.add_paragraph()

# Información del documento
info_table = doc.add_table(rows=6, cols=2)
info_table.style = 'Light Grid Accent 1'
info_data = [
    ('Proyecto:', 'Facturacion Automatizada ICA - SNRI + SIGMA'),
    ('Modulo:', 'Extraccion OCR de Comprobantes de Recaudo'),
    ('Version:', '1.0'),
    ('Fecha:', datetime.date.today().strftime('%d/%m/%Y')),
    ('Elaborado por:', 'Equipo Tecnico Facturacion ICA'),
    ('Destinatarios:', 'Direccion Financiera / Oficina de Tecnologias de la Informacion (OTI)'),
]
for i, (label, value) in enumerate(info_data):
    info_table.cell(i, 0).text = label
    info_table.cell(i, 1).text = value
    for cell in info_table.rows[i].cells:
        for paragraph in cell.paragraphs:
            paragraph.style = doc.styles['Normal']
            for run in paragraph.runs:
                run.font.size = Pt(10)

doc.add_paragraph()

# 1. CONTEXTO
doc.add_heading('1. CONTEXTO Y OBJETIVO', level=2)
doc.add_paragraph(
    'El Instituto Colombiano Agropecuario (ICA) esta implementando un sistema de facturacion automatizada '
    'que integra el Sistema Nacional de Recaudo (SNRI) con el Sistema de Informacion para Guias de '
    'Movilizacion Animal (SIGMA). Como parte de este proyecto, se ha desarrollado un modulo de '
    'reconocimiento optico de caracteres (OCR) para extraer automaticamente los datos clave de los '
    'comprobantes de pago (vouchers) que presentan los usuarios al realizar pagos en las diferentes '
    'entidades bancarias conveniadas.'
)
doc.add_paragraph(
    'El objetivo de esta solicitud es obtener una muestra representativa de comprobantes reales para:'
)

bullets = [
    'Entrenar y ajustar los modelos de OCR (Tesseract y PaddleOCR) con documentos reales del ICA',
    'Validar y refinar las expresiones regulares (regex) de los 6 perfiles de extraccion configurados',
    'Medir la precision de extraccion (accuracy) por tipo de comprobante y canal bancario',
    'Identificar variaciones de formato, calidad de imagen, y casos edge no contemplados',
    'Garantizar que el sistema funcione en produccion desde el dia 1 con la mayor precision posible'
]
for b in bullets:
    doc.add_paragraph(b, style='List Bullet')

doc.add_paragraph()

# 2. PERFILES DE EXTRACCIÓN CONFIGURADOS
doc.add_heading('2. PERFILES DE EXTRACCIÓN CONFIGURADOS (GRF-GI-GUIA-002)', level=2)
doc.add_paragraph(
    'Basados en la Guia GRF-GI-GUIA-002 V2 "Guia para el Registro de Numero de Soporte de Recaudo", '
    'se han configurado los siguientes 6 perfiles de extraccion que cubren los 10 canales bancarios '
    'identificados:'
)

# Tabla de perfiles
perfil_table = doc.add_table(rows=7, cols=4)
perfil_table.style = 'Light Grid Accent 1'
perfil_table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['Codigo', 'Nombre / Banco', 'Canales Cubiertos', 'Campos Clave a Extraer']
for i, h in enumerate(headers):
    cell = perfil_table.cell(0, i)
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(9)

perfiles_data = [
    ('davivienda', 'Davivienda', 'Consignacion - Corresponsales - Punto de Pago', 'Fecha, Valor, N Transaccion, Convenio, Entidad ICA'),
    ('colpatria', 'Colpatria', 'Punto de Pago', 'Fecha, Valor, Id Trx, PIN ICA, Convenio, Entidad ICA'),
    ('punto_red', 'Punto Red', 'Punto Red', 'Fecha, Valor, N Transaccion, PIN ICA, Entidad ICA'),
    ('bancolombia', 'Bancolombia', 'Sucursal - Corresponsal - Cajero Automatico', 'Fecha, Valor, N Transaccion, Registro Operacion, TRANS, Convenio, Entidad ICA'),
    ('banco_occidente', 'Banco de Occidente', 'Consignacion', 'Fecha, Valor, N Transaccion, Entidad ICA, Cuenta (ultimos 4), Convenio'),
    ('datafono', 'Datafono / POS Generico', 'Datafono (POS)', 'Fecha, Valor, AUT/Autorizacion, Forma Pago=DATAFONO'),
]

for row_idx, (codigo, nombre, canales, campos) in enumerate(perfiles_data, 1):
    perfil_table.cell(row_idx, 0).text = codigo
    perfil_table.cell(row_idx, 1).text = nombre
    perfil_table.cell(row_idx, 2).text = canales
    perfil_table.cell(row_idx, 3).text = campos
    for cell in perfil_table.rows[row_idx].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)

doc.add_paragraph()

# 3. SOLICITUD ESPECÍFICA DE MUESTRAS
doc.add_heading('3. SOLICITUD ESPECIFICA DE MUESTRAS', level=2)

doc.add_heading('3.1 Cantidad Minima Requerida', level=3)
cant_table = doc.add_table(rows=7, cols=3)
cant_table.style = 'Light Grid Accent 1'
cant_headers = ['Perfil / Banco', 'Minimo Requerido', 'Ideal (para validacion robusta)']
for i, h in enumerate(cant_headers):
    cant_table.cell(0, i).text = h
    for p in cant_table.cell(0, i).paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)

cant_data = [
    ('Davivienda (3 canales)', '15 (5 c/u)', '30 (10 c/u)'),
    ('Colpatria Punto de Pago', '10', '20'),
    ('Punto Red', '10', '20'),
    ('Bancolombia (3 canales)', '15 (5 c/u)', '30 (10 c/u)'),
    ('Banco de Occidente', '10', '20'),
    ('Datafono / POS', '10', '20'),
]
for row_idx, (perfil, min_req, ideal) in enumerate(cant_data, 1):
    cant_table.cell(row_idx, 0).text = perfil
    cant_table.cell(row_idx, 1).text = min_req
    cant_table.cell(row_idx, 2).text = ideal
    for cell in cant_table.rows[row_idx].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

doc.add_paragraph()
p = doc.add_paragraph('Total minimo: 70 comprobantes  |  Total ideal: 140 comprobantes')
p.runs[0].bold = True

doc.add_heading('3.2 Caracteristicas de las Muestras', level=3)
reqs = [
    'Formato: Imagen (JPG/PNG) o PDF - tal como los recibe el usuario',
    'Calidad: Legibles, sin recortes, con toda la informacion visible',
    'Variedad: Incluir diferentes calidades (buena, regular, baja), orientaciones, y tipos de impresora',
    'Datos reales: Deben corresponder a transacciones reales (pueden anonimizar NIT/valores si es necesario)',
    'Nombrenado sugerido: banco_canal_YYYYMMDD_XXX.jpg (ej: davivienda_consignacion_20260115_001.jpg)',
    'Entrega: Carpeta compartida (OneDrive/SharePoint) o medio seguro definido por OTI',
]
for r in reqs:
    doc.add_paragraph(r, style='List Bullet')

doc.add_heading('3.3 Casos Especiales de Interes (Prioridad Alta)', level=3)
casos = [
    'Comprobantes con ceros a la izquierda en el N de transaccion (ej: "08757830" en Banco Occidente)',
    'Comprobantes con letras en el N de transaccion',
    'Imagenes rotadas, inclinadas, o con ruido/fondo complejo',
    'Comprobantes de cajero automatico Bancolombia con referencia a factura SNRI',
    'Datafonos con diferentes formatos de AUT/Autorizacion',
    'Corresponsales bancarios con formatos de "Aprob Banco" o "Apro"',
    'Punto Red / Colpatria con PIN ICA visible',
]
for c in casos:
    doc.add_paragraph(c, style='List Bullet')

doc.add_paragraph()

# 4. CAMPOS QUE SE EXTRAEN AUTOMÁTICAMENTE
doc.add_heading('4. CAMPOS QUE EL SISTEMA EXTRAE AUTOMATICAMENTE', level=2)
doc.add_paragraph(
    'Para cada comprobante, el OCR extrae los siguientes campos (configurables como obligatorios/opcionales):'
)

campos_table = doc.add_table(rows=18, cols=3)
campos_table.style = 'Light Grid Accent 1'
campos_headers = ['Campo Interno', 'Etiqueta Visible', 'Descripcion / Ejemplo']
for i, h in enumerate(campos_headers):
    campos_table.cell(0, i).text = h
    for p in campos_table.cell(0, i).paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)

campos_data = [
    ('fecha', 'Fecha', 'Fecha de la transaccion (DD/MM/YYYY)'),
    ('valor_total', 'Valor Total', 'Monto pagado (formato colombiano: 1.500.000)'),
    ('numero_transaccion', 'N Transaccion', 'Transaccion / Aprob Banco / Id Trx / AUT / Registro Operacion / TRANS'),
    ('aprob_banco', 'Aprob. Banco', 'Especifico Davivienda Corresponsales'),
    ('id_trx', 'Id Trx', 'Especifico Davivienda Punto Pago / Colpatria'),
    ('registro_operacion', 'Registro Operacion', 'Especifico Bancolombia Sucursal'),
    ('trans_cajero', 'TRANS Cajero', 'Especifico Bancolombia Cajero Automatico'),
    ('convenio_recaudo', 'Convenio Recaudo', 'N de convenio bancario / PIN ICA'),
    ('entidad_recaudo', 'Entidad Recaudo', 'INSTITUTO COLOMBIANO AGROPECUARIO / ICA RECAUDOS'),
    ('cuenta_ultimos4', 'Cuenta (ultimos 4)', 'Ultimos 4 digitos de cuenta (Banco Occidente)'),
    ('pin_ica', 'PIN ICA', 'Codigo PIN ICA (Colpatria / Punto Red)'),
    ('banco', 'Banco', 'Nombre del banco (literal segun perfil)'),
    ('forma_pago', 'Forma de Pago', 'CONSIGNACION / DATAFONO / TRANSFERENCIA / etc.'),
    ('numero_referencia', 'N Referencia', 'N de factura SNRI o referencia de pago'),
    ('convenio', 'Convenio', 'Convenio generico (compatibilidad)'),
    ('ubicacion', 'Ubicacion', 'Sucursal / corresponsal / punto de pago'),
    ('autenticacion', 'Autenticacion', 'Codigo de autorizacion (DATAFONO)'),
]

for row_idx, (campo, etiqueta, desc) in enumerate(campos_data, 1):
    campos_table.cell(row_idx, 0).text = campo
    campos_table.cell(row_idx, 1).text = etiqueta
    campos_table.cell(row_idx, 2).text = desc
    for cell in campos_table.rows[row_idx].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

doc.add_paragraph()

# 5. BENEFICIOS ESPERADOS
doc.add_heading('5. BENEFICIOS ESPERADOS DEL SISTEMA', level=2)
beneficios = [
    'Reduccion del 80-90% del tiempo de digitacion manual de comprobantes',
    'Eliminacion de errores de transcripcion (NIT, valores, referencias)',
    'Trazabilidad completa: comprobante - OCR - validacion - factura SNRI - sync SIGMA',
    'Cumplimiento de la Guia GRF-GI-GUIA-002 y normativa SNRI/DIAN',
    'Dashboard de metricas en tiempo real: comprobantes procesados, tasa de exito, confianza OCR',
    'Re-extraccion con cambio de perfil sin repetir OCR (ahorro de tiempo)',
    'Modo demo para pruebas sin certificado SNRI real',
]
for b in beneficios:
    doc.add_paragraph(b, style='List Bullet')

doc.add_paragraph()

# 6. CRONOGRAMA Y ENTREGABLES
doc.add_heading('6. CRONOGRAMA PROPUESTO', level=2)
crono_table = doc.add_table(rows=6, cols=3)
crono_table.style = 'Light Grid Accent 1'
crono_headers = ['Fase', 'Actividad', 'Plazo Estimado']
for i, h in enumerate(crono_headers):
    crono_table.cell(0, i).text = h
    for p in crono_table.cell(0, i).paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)

crono_data = [
    ('1. Recoleccion', 'Entrega de muestras por Financiera/OTI', 'Semana 1-2'),
    ('2. Entrenamiento', 'Ajuste de regex + validacion PaddleOCR', 'Semana 2-3'),
    ('3. Pruebas', 'Pruebas de precision con conjunto de validacion', 'Semana 3-4'),
    ('4. Piloto', 'Piloto controlado con usuarios reales', 'Semana 4-6'),
    ('5. Produccion', 'Despliegue productivo completo', 'Semana 6+'),
]

for row_idx, (fase, act, plazo) in enumerate(crono_data, 1):
    crono_table.cell(row_idx, 0).text = fase
    crono_table.cell(row_idx, 1).text = act
    crono_table.cell(row_idx, 2).text = plazo
    for cell in crono_table.rows[row_idx].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

doc.add_paragraph()

# 7. CONTACTO Y RESPONSABLES
doc.add_heading('7. CONTACTO Y RESPONSABLES', level=2)
contact_table = doc.add_table(rows=4, cols=3)
contact_table.style = 'Light Grid Accent 1'
contact_headers = ['Rol', 'Nombre / Area', 'Contacto']
for i, h in enumerate(contact_headers):
    contact_table.cell(0, i).text = h
    for p in contact_table.cell(0, i).paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)

contact_data = [
    ('Lider Tecnico OCR', '[Nombre] - Equipo Facturacion', '[email@ica.gov.co] / [ext.]'),
    ('Coordinador Financiera', '[Nombre] - Direccion Financiera', '[email@ica.gov.co] / [ext.]'),
    ('Responsable OTI', '[Nombre] - Oficina Tec. Informacion', '[email@ica.gov.co] / [ext.]'),
]

for row_idx, (rol, nombre, contacto) in enumerate(contact_data, 1):
    contact_table.cell(row_idx, 0).text = rol
    contact_table.cell(row_idx, 1).text = nombre
    contact_table.cell(row_idx, 2).text = contacto
    for cell in contact_table.rows[row_idx].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

doc.add_paragraph()

# 8. ANEXOS
doc.add_heading('8. ANEXOS DE REFERENCIA', level=2)
anexos = [
    'Guia GRF-GI-GUIA-002 V2 - Guia para el Registro de Numero de Soporte de Recaudo',
    'Manual de Identidad Visual ICA (colores institucionales aplicados en la interfaz)',
    'Esquema de base de datos: tabla perfiles_extraccion, comprobantes, facturas',
    'API REST: POST /api/v1/upload, GET /api/v1/perfiles, POST /api/v1/perfiles/seed',
    'Documentacion tecnica: README.md, PLAN_IMPLEMENTACION.md, VALIDACION_PADDLEOCR.md',
]
for a in anexos:
    doc.add_paragraph(a, style='List Bullet')

doc.add_paragraph()

# Firma
doc.add_paragraph('_' * 60)
doc.add_paragraph('Aprobado por:').runs[0].bold = True
doc.add_paragraph('_________________________________')
doc.add_paragraph('Direccion Financiera - ICA')
doc.add_paragraph()
doc.add_paragraph('_________________________________')
doc.add_paragraph('Oficina de Tecnologias de la Informacion (OTI) - ICA')

# Guardar
output_path = 'Solicitud_Muestras_Vouchers_OCR_ICA.docx'
doc.save(output_path)
print(f'Documento guardado: {output_path}')